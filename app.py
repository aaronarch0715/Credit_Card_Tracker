"""
신용카드 명세서 분석기 (Streamlit)

카드사 엑셀 명세서를 업로드하면 가맹점명 키워드로 사용처를 자동 분류하고,
사용자가 표에서 분류를 직접 고친 뒤 대시보드와 엑셀 결과물을 받을 수 있다.
사용자는 분류 키워드와 새 분류(색상 포함)를 직접 추가할 수 있다.
"""

from __future__ import annotations

import hashlib
import io
import json
import re
from pathlib import Path
from typing import Optional

import pandas as pd
import plotly.express as px
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
import streamlit as st

# ─────────────────────────────────────────────────────────────
# 1. 분류 사전 & 상수
# ─────────────────────────────────────────────────────────────

# 순서가 중요하다: 위에 있는 카테고리가 먼저 매칭된다.
# (예: '카카오T'는 교통으로, '네이버페이'는 쇼핑으로 가야 하므로 구체적인 키워드를 앞에 둔다)
CATEGORY_KEYWORDS: dict[str, list[str]] = {
    "교통/차량": [
        "카카오T", "카카오택시", "카카오모빌리티", "택시", "지하철", "버스", "교통",
        "티머니", "캐시비", "코레일", "KTX", "SRT", "에스알", "고속버스", "시외버스",
        "주유", "오일", "SK에너지", "GS칼텍스", "S-OIL", "에쓰오일", "현대오일뱅크",
        "알뜰주유", "하이패스", "도로공사", "톨게이트", "주차", "쏘카", "그린카",
        "타다", "우버", "UBER", "항공", "대한항공", "아시아나", "제주항공", "진에어",
        "티웨이", "세차", "자동차", "타이어",
    ],
    "식비/카페": [
        "스타벅스", "STARBUCKS", "투썸", "이디야", "메가커피", "메가엠지씨", "컴포즈",
        "빽다방", "폴바셋", "할리스", "커피빈", "블루보틀", "카페", "커피", "CAFE",
        "COFFEE", "배달의민족", "배민", "우아한형제들", "요기요", "쿠팡이츠",
        "맥도날드", "버거킹", "롯데리아", "KFC", "맘스터치", "서브웨이", "써브웨이",
        "파리바게뜨", "파리바게트", "뚜레쥬르", "베이커리", "제과", "던킨", "배스킨",
        "베스킨", "설빙", "도미노", "피자", "치킨", "BBQ", "BHC", "교촌", "굽네",
        "식당", "음식", "분식", "김밥", "국밥", "한식", "중식", "일식", "양식",
        "레스토랑", "고기", "갈비", "삼겹", "족발", "보쌈", "초밥", "스시", "라멘",
        "푸드", "FOOD", "다이닝", "주점", "호프", "포차",
    ],
    "쇼핑": [
        "쿠팡", "COUPANG", "네이버페이", "NAVERPAY", "네이버", "11번가", "G마켓",
        "지마켓", "옥션", "SSG", "쓱", "위메프", "티몬", "마켓컬리", "컬리", "오늘의집",
        "29CM", "지그재그", "에이블리", "알리익스프레스", "ALIEXPRESS", "테무", "TEMU",
        "아마존", "AMAZON", "올리브영", "무신사", "MUSINSA", "다이소", "이마트",
        "홈플러스", "롯데마트", "코스트코", "COSTCO", "트레이더스", "마트", "백화점",
        "아울렛", "편의점", "GS25", "CU", "세븐일레븐", "이마트24", "미니스톱",
        "유니클로", "자라", "ZARA", "H&M", "탑텐", "스파오", "ABC마트", "다이슨",
        "하이마트", "전자랜드", "애플", "APPLE", "삼성스토어",
    ],
    "주거/통신": [
        "SKT", "SK텔레콤", "KT", "케이티", "LGU", "LG유플러스", "LG U+", "유플러스",
        "알뜰폰", "통신", "인터넷", "관리비", "아파트", "도시가스", "가스", "한전",
        "한국전력", "전기", "수도", "월세", "부동산", "SK브로드밴드", "브로드밴드",
    ],
    "의료/건강": [
        "병원", "의원", "약국", "치과", "한의원", "안과", "피부과", "내과", "외과",
        "소아과", "이비인후과", "정형외과", "산부인과", "정신건강", "검진", "메디컬",
        "클리닉", "헬스", "필라테스", "요가", "피트니스", "GYM",
    ],
    "문화/여가": [
        "CGV", "메가박스", "롯데시네마", "영화", "넷플릭스", "NETFLIX", "유튜브",
        "YOUTUBE", "디즈니", "DISNEY", "왓챠", "티빙", "TVING", "웨이브", "WAVVE",
        "쿠팡플레이", "스포티파이", "SPOTIFY", "멜론", "지니뮤직", "벅스", "인터파크",
        "INTERPARK", "예스24", "YES24", "교보문고", "알라딘", "영풍문고", "서점",
        "티켓", "공연", "전시", "박물관", "미술관", "놀이공원", "에버랜드", "롯데월드",
        "노래", "코인노래", "PC방", "볼링", "골프", "스크린", "게임", "STEAM",
        "닌텐도", "NINTENDO", "플레이스테이션", "호텔", "모텔", "야놀자", "여기어때",
        "에어비앤비", "AIRBNB", "아고다", "AGODA", "익스피디아", "여행",
    ],
}

DEFAULT_CATEGORY = "기타"
# 기본 분류의 화면/차트 표시 순서 (사용자 분류는 '기타' 바로 앞에 붙는다)
BUILTIN_ORDER = ["식비/카페", "교통/차량", "쇼핑", "주거/통신", "의료/건강", "문화/여가"]

# 기본 분류 색상 (고정). 순위가 아닌 '카테고리'를 따라가므로 필터가 바뀌어도 색이 유지된다.
BUILTIN_COLORS = {
    "식비/카페": "#2a78d6",
    "교통/차량": "#eb6834",
    "쇼핑": "#1baf7a",
    "주거/통신": "#eda100",
    "의료/건강": "#e87ba4",
    "문화/여가": "#6250d6",
    DEFAULT_CATEGORY: "#9a9993",
}
# 사용자가 새 분류를 만들 때 추천하는 색상 (기본 분류 색과 겹치지 않는 것부터)
SUGGESTED_COLORS = ["#008300", "#e34948", "#0f8b8d", "#8a5a00", "#b5179e", "#4d7c0f", "#5c6bc0", "#d4a017"]

# 열 자동 인식용 키워드 (카드사별 표기 차이 대응, 앞쪽일수록 우선)
COLUMN_HINTS = {
    "merchant": ["가맹점명", "이용가맹점", "가맹점", "이용하신곳", "이용처", "사용처", "상호", "거래처", "적요", "내용"],
    "amount": ["결제금액", "이용금액", "거래금액", "사용금액", "청구금액", "승인금액", "결제원금", "금액"],
    "date": ["이용일자", "거래일자", "이용일", "거래일", "사용일자", "사용일", "승인일자", "승인일", "일자", "날짜"],
}

# 화면에 표시되는 고정 열 구성: 번호 · 이용일자 · 가맹점 · 결제금액 · 분류
COL_NO, COL_DATE, COL_MERCHANT, COL_AMOUNT, COL_CATEGORY = "번호", "이용일자", "가맹점", "결제금액", "분류"
COL_RID = "_rid"  # 행 고유번호 (필터가 바뀌어도 수동 분류가 유지되도록 내부에서만 사용)

# 사용자가 추가한 분류/키워드 저장 파일
SETTINGS_PATH = Path(__file__).with_name("user_keywords.json")


# ─────────────────────────────────────────────────────────────
# 2. 사용자 설정 (내 키워드 / 내 분류)
# ─────────────────────────────────────────────────────────────

def load_settings() -> dict:
    """{"categories": [{"name", "color"}], "keywords": {분류: [키워드, ...]}}"""
    try:
        data = json.loads(SETTINGS_PATH.read_text(encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError, OSError):
        data = {}
    categories = [
        {"name": str(c["name"]), "color": str(c.get("color") or SUGGESTED_COLORS[0])}
        for c in data.get("categories", [])
        if isinstance(c, dict) and c.get("name") and c["name"] not in BUILTIN_COLORS
    ]
    keywords = {
        str(cat): [str(k) for k in kws if str(k).strip()]
        for cat, kws in (data.get("keywords") or {}).items()
        if isinstance(kws, list)
    }
    return {"categories": categories, "keywords": keywords}


def save_settings(settings: dict) -> None:
    try:
        SETTINGS_PATH.write_text(json.dumps(settings, ensure_ascii=False, indent=2), encoding="utf-8")
    except OSError:
        st.session_state["flash"] = ("warning", "설정 파일을 저장하지 못했어요. 이번 접속 동안에만 적용됩니다.")


def get_settings() -> dict:
    if "settings" not in st.session_state:
        st.session_state["settings"] = load_settings()
    return st.session_state["settings"]


def user_category_names(settings: dict) -> list[str]:
    return [c["name"] for c in settings["categories"]]


def category_order(settings: dict) -> list[str]:
    return BUILTIN_ORDER + user_category_names(settings) + [DEFAULT_CATEGORY]


def category_colors(settings: dict) -> dict[str, str]:
    return {**BUILTIN_COLORS, **{c["name"]: c["color"] for c in settings["categories"]}}


def _reset_keyword_widgets() -> None:
    """저장된 설정이 바뀌면 사이드바 위젯 상태를 지워 새 값으로 다시 그리게 한다."""
    for key in list(st.session_state.keys()):
        if key.startswith(("kw_", "color_")):
            del st.session_state[key]


def on_keywords_change(category: str) -> None:
    settings = get_settings()
    new_list: list[str] = []
    for kw in st.session_state.get(f"kw_{category}", []):
        kw = str(kw).strip()
        if kw and _normalize(kw) not in {_normalize(k) for k in new_list}:
            new_list.append(kw)
    # 같은 키워드가 다른 분류에도 있으면 그쪽에서 빼서 '마지막으로 지정한 분류'가 이기게 한다
    added = {_normalize(k) for k in new_list}
    for other, kws in settings["keywords"].items():
        if other != category:
            settings["keywords"][other] = [k for k in kws if _normalize(k) not in added]
    settings["keywords"][category] = new_list
    save_settings(settings)
    _reset_keyword_widgets()


def on_color_change(category: str) -> None:
    settings = get_settings()
    for c in settings["categories"]:
        if c["name"] == category:
            c["color"] = st.session_state[f"color_{category}"]
    save_settings(settings)


def on_delete_category(category: str) -> None:
    settings = get_settings()
    settings["categories"] = [c for c in settings["categories"] if c["name"] != category]
    settings["keywords"].pop(category, None)
    save_settings(settings)
    _reset_keyword_widgets()
    st.session_state["flash"] = ("success", f"'{category}' 분류를 삭제했어요. 해당 내역은 다시 자동 분류됩니다.")


def on_add_category() -> None:
    settings = get_settings()
    name = st.session_state.get("new_cat_name", "").strip()
    color = st.session_state.get("new_cat_color", SUGGESTED_COLORS[0])
    raw_kws = st.session_state.get("new_cat_keywords", "")
    if not name:
        st.session_state["flash"] = ("warning", "분류 이름을 입력해 주세요.")
        return
    if name in category_order(settings):
        st.session_state["flash"] = ("warning", f"'{name}' 분류는 이미 있어요.")
        return
    settings["categories"].append({"name": name, "color": color})
    save_settings(settings)
    st.session_state[f"kw_{name}"] = [k.strip() for k in raw_kws.split(",") if k.strip()]
    on_keywords_change(name)  # 키워드 저장 + 다른 분류와의 중복 정리
    st.session_state["flash"] = ("success", f"'{name}' 분류를 추가했어요.")


def next_suggested_color(settings: dict) -> str:
    used = {c.lower() for c in category_colors(settings).values()}
    return next((c for c in SUGGESTED_COLORS if c.lower() not in used), SUGGESTED_COLORS[0])


# ─────────────────────────────────────────────────────────────
# 3. 데이터 처리 함수
# ─────────────────────────────────────────────────────────────

def _normalize(text: str) -> str:
    """키워드 비교용: 공백 제거 + 대문자화"""
    return re.sub(r"\s+", "", str(text)).upper()


# 정규화된 기본 키워드를 미리 계산해 둔다
_BUILTIN_RULES = [(cat, [_normalize(k) for k in kws]) for cat, kws in CATEGORY_KEYWORDS.items()]


def build_user_rules(settings: dict) -> list[tuple[str, str]]:
    """사용자 키워드 → (정규화 키워드, 분류). 긴 키워드가 먼저 매칭되도록 정렬."""
    valid = set(category_order(settings))
    rules = [
        (_normalize(k), cat)
        for cat, kws in settings["keywords"].items() if cat in valid
        for k in kws if _normalize(k)
    ]
    return sorted(rules, key=lambda r: -len(r[0]))


def classify_merchant(name: str, user_rules: list[tuple[str, str]]) -> str:
    """가맹점명으로 분류를 추정한다. 사용자 키워드가 기본 키워드보다 우선한다."""
    if not isinstance(name, str) or not name.strip():
        return DEFAULT_CATEGORY
    target = _normalize(name)
    for kw, cat in user_rules:
        if kw in target:
            return cat
    for cat, kws in _BUILTIN_RULES:
        if any(k in target for k in kws):
            return cat
    return DEFAULT_CATEGORY


def parse_amount(value) -> Optional[float]:
    """'12,300원', '₩ 5,000', '(1,000)', '-3,000' 등을 숫자로 변환. 실패하면 None."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip()
    if not s:
        return None
    negative = s.startswith("-") or (s.startswith("(") and s.endswith(")"))
    digits = re.sub(r"[^0-9.]", "", s)
    if not digits or digits == ".":
        return None
    try:
        num = float(digits)
    except ValueError:
        return None
    return -num if negative else num


def parse_dates(series: pd.Series) -> pd.Series:
    """'2024.01.05', '2024/01/05', '20240105', 엑셀 날짜 등 다양한 형식을 datetime으로."""
    if pd.api.types.is_datetime64_any_dtype(series):
        return series
    s = series.map(lambda v: "" if pd.isna(v) else str(v)).astype(object).str.strip()
    s = s.str.replace(r"\s+\d{1,2}:\d{2}(:\d{2})?$", "", regex=True)  # 뒤에 붙은 시간 제거
    s = s.str.replace(r"[./]", "-", regex=True).str.rstrip("-")
    # 20240105 같은 8자리 숫자
    s = s.where(~s.str.fullmatch(r"\d{8}", na=False), s.str.slice(0, 4) + "-" + s.str.slice(4, 6) + "-" + s.str.slice(6, 8))
    # 26-10-03 같은 2자리 연도
    s = s.where(~s.str.fullmatch(r"\d{2}-\d{1,2}-\d{1,2}", na=False), "20" + s)
    # 10-03 처럼 연도가 없는 경우: 같은 열의 다른 날짜에서 연도를 추정 (없으면 올해)
    no_year = s.str.fullmatch(r"\d{1,2}-\d{1,2}", na=False)
    if no_year.any():
        years = s[~no_year].str.extract(r"^(\d{4})-", expand=False).dropna()
        year = years.mode().iloc[0] if not years.empty else str(pd.Timestamp.today().year)
        s = s.where(~no_year, year + "-" + s)
    return pd.to_datetime(s, errors="coerce", format="mixed")


def suggest_column(columns: list[str], kind: str) -> Optional[str]:
    """열 이름 키워드 매칭으로 해당 열을 찾는다. 앞쪽 힌트일수록 우선순위가 높다."""
    norm_cols = {c: _normalize(c) for c in columns}
    for hint in COLUMN_HINTS[kind]:
        h = _normalize(hint)
        for col, nc in norm_cols.items():  # 정확히 일치
            if nc == h:
                return col
    for hint in COLUMN_HINTS[kind]:
        h = _normalize(hint)
        for col, nc in norm_cols.items():  # 부분 일치
            if h in nc:
                return col
    return None


def detect_header_row(raw: pd.DataFrame, max_scan: int = 30) -> int:
    """
    카드사 엑셀은 상단에 제목/고객정보 행이 있는 경우가 많다.
    앞쪽 행들 중 열 이름 키워드가 가장 많이 들어 있는 행을 헤더로 판단한다.
    """
    all_hints = [_normalize(h) for hints in COLUMN_HINTS.values() for h in hints]
    best_row, best_score = 0, 0
    for i in range(min(max_scan, len(raw))):
        cells = [_normalize(v) for v in raw.iloc[i].tolist() if pd.notna(v)]
        score = sum(any(h in c for h in all_hints) for c in cells)
        # 점수가 같으면 아래 행을 택한다 (2줄짜리 병합 헤더에서는 아래 줄이 실제 열 이름)
        if score >= best_score and score > 0:
            best_row, best_score = i, score
    return best_row


def detect_file_format(file_bytes: bytes) -> str:
    """
    확장자가 아니라 파일 내용으로 실제 형식을 판별한다.
    국내 카드사의 '.xls' 다운로드는 실제로는 HTML 표나 XML(SpreadsheetML)인 경우가 많다.
    """
    head = file_bytes[:8]
    if head.startswith(b"PK"):
        return "xlsx"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "xls"
    text = _decode_text(file_bytes[:4096]).lower()
    if "urn:schemas-microsoft-com:office:spreadsheet" in text and "<workbook" in text:
        return "xml"
    if "<table" in text or "<html" in text or "<!doctype" in text or "<meta" in text:
        return "html"
    return "unknown"


def _decode_text(data: bytes) -> str:
    """BOM/UTF-8/CP949(EUC-KR) 순으로 시도해 텍스트로 디코딩"""
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8", errors="replace")
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace")
    for enc in ("utf-8", "cp949"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("cp949", errors="replace")


def _grid_to_frame(rows: list[list]) -> pd.DataFrame:
    width = max((len(r) for r in rows), default=0)
    return pd.DataFrame([r + [None] * (width - len(r)) for r in rows], dtype=object)


def _span(cell, attr: str) -> int:
    value = str(cell.get(attr, "1")).strip()
    return max(1, min(int(value), 100)) if value.isdigit() else 1


def _read_html_tables(file_bytes: bytes) -> dict[str, pd.DataFrame]:
    """HTML 형식 '가짜 xls'의 표들을 읽는다. 숫자/날짜 자동 변환 없이 문자열 그대로 가져온다."""
    import lxml.html

    root = lxml.html.fromstring(_decode_text(file_bytes))
    tables = root.xpath("//table")
    # 레이아웃용 바깥 표는 빼고, 안쪽에 다른 표가 없는 '실제 데이터 표'만 사용
    leaf_tables = [t for t in tables if not t.xpath(".//table")] or tables

    result: dict[str, pd.DataFrame] = {}
    for idx, table in enumerate(leaf_tables, start=1):
        grid: list[list] = []
        pending: dict[int, tuple[int, str]] = {}  # rowspan 처리: 열 → (남은 행 수, 값)
        for tr in table.xpath("./tr|./thead/tr|./tbody/tr|./tfoot/tr"):
            row: list = []
            cells = iter(tr.xpath("./td|./th"))
            col = 0
            while True:
                if col in pending:  # 위 행에서 rowspan으로 내려온 칸
                    left, val = pending[col]
                    row.append(val)
                    pending[col] = (left - 1, val)
                    if left - 1 <= 0:
                        del pending[col]
                    col += 1
                    continue
                cell = next(cells, None)
                if cell is None:
                    break
                text = re.sub(r"\s+", " ", cell.text_content()).strip() or None
                colspan, rowspan = _span(cell, "colspan"), _span(cell, "rowspan")
                for k in range(colspan):
                    row.append(text if k == 0 else None)
                    if rowspan > 1:
                        pending[col] = (rowspan - 1, text if k == 0 else None)
                    col += 1
            if any(v is not None for v in row):
                grid.append(row)
        if len(grid) >= 2 and max(len(r) for r in grid) >= 2:
            result[f"표{idx}"] = _grid_to_frame(grid)

    # 거래내역 표가 보통 가장 크므로 행 수가 많은 표를 먼저 본다
    return dict(sorted(result.items(), key=lambda kv: -len(kv[1])))


def _read_spreadsheet_xml(file_bytes: bytes) -> dict[str, pd.DataFrame]:
    """Excel 2003 XML(SpreadsheetML) 형식을 읽는다."""
    import xml.etree.ElementTree as ET

    ns = "{urn:schemas-microsoft-com:office:spreadsheet}"
    root = ET.fromstring(file_bytes)
    result: dict[str, pd.DataFrame] = {}
    for i, ws in enumerate(root.iter(f"{ns}Worksheet"), start=1):
        name = ws.get(f"{ns}Name") or f"Sheet{i}"
        grid: list[list] = []
        for row in ws.iter(f"{ns}Row"):
            values: list = []
            for cell in row.iter(f"{ns}Cell"):
                pos = cell.get(f"{ns}Index")
                if pos and pos.isdigit():  # 빈 칸을 건너뛴 경우
                    values += [None] * (int(pos) - 1 - len(values))
                data = cell.find(f"{ns}Data")
                values.append("".join(data.itertext()).strip() if data is not None else None)
                values += [None] * int(cell.get(f"{ns}MergeAcross", "0") or 0)
            grid.append(values)
        if grid:
            result[name] = _grid_to_frame(grid)
    return result


@st.cache_data(show_spinner=False)
def read_raw_excel(file_bytes: bytes, file_name: str) -> dict[str, pd.DataFrame]:
    """엑셀의 모든 시트를 header 없이 읽는다 (헤더 행은 나중에 결정)."""
    fmt = detect_file_format(file_bytes)
    if fmt == "html":
        return _read_html_tables(file_bytes)
    if fmt == "xml":
        return _read_spreadsheet_xml(file_bytes)
    if fmt == "unknown":  # 시그니처로 판별 불가 → 확장자 기준으로 시도
        fmt = "xls" if file_name.lower().endswith(".xls") else "xlsx"
    engine = "xlrd" if fmt == "xls" else "openpyxl"
    return pd.read_excel(io.BytesIO(file_bytes), sheet_name=None, header=None, engine=engine, dtype=object)


def apply_header(raw: pd.DataFrame, header_row: int) -> pd.DataFrame:
    """지정한 행을 열 이름으로 사용하고 그 아래를 데이터로 만든다."""
    header = raw.iloc[header_row].tolist()
    cols, seen = [], {}
    for i, h in enumerate(header):
        name = str(h).strip() if pd.notna(h) and str(h).strip() else f"열{i + 1}"
        if name in seen:  # 중복 열 이름 처리
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 0
        cols.append(name)
    df = raw.iloc[header_row + 1:].copy()
    df.columns = cols
    return df.dropna(how="all").reset_index(drop=True)


def locate_statement(sheets: dict[str, pd.DataFrame]):
    """
    시트(표)들 중 가맹점·결제금액 열을 찾을 수 있는 첫 번째 표를 고른다.
    반환: (시트명, 데이터프레임, {"date": 열, "merchant": 열, "amount": 열}) 또는 None
    """
    for name, raw in sheets.items():
        df = apply_header(raw, detect_header_row(raw))
        cols = list(df.columns)
        found = {kind: suggest_column(cols, kind) for kind in ("date", "merchant", "amount")}
        if found["merchant"] and found["amount"] and found["merchant"] != found["amount"]:
            if found["date"] in (found["merchant"], found["amount"]):
                found["date"] = None
            return name, df, found
    return None


def build_transactions(df: pd.DataFrame, merchant_col: str, amount_col: str, date_col: Optional[str]) -> pd.DataFrame:
    """찾아낸 열로 표준 거래내역 표(이용일자·가맹점·결제금액)를 만든다."""
    out = pd.DataFrame()
    out[COL_DATE] = parse_dates(df[date_col]) if date_col else pd.Series(pd.NaT, index=df.index)
    out[COL_MERCHANT] = df[merchant_col].map(lambda v: "" if pd.isna(v) else str(v).strip())
    out[COL_AMOUNT] = df[amount_col].map(parse_amount)

    # 금액이 없거나 가맹점이 비어 있는 행(소계/합계/빈 줄 등) 제거
    out = out[out[COL_AMOUNT].notna()]
    out = out[~out[COL_MERCHANT].isin(["", "nan", "None"])]
    out = out[~out[COL_MERCHANT].str.contains(r"^(?:합계|소계|총계|총\s*합계)", regex=True, na=False)]

    out[COL_AMOUNT] = out[COL_AMOUNT].round(0).astype("int64")
    out = out.reset_index(drop=True)
    out[COL_RID] = range(len(out))
    return out


def summarize(df: pd.DataFrame, order: list[str]) -> pd.DataFrame:
    """분류별 합계/건수/비중"""
    g = df.groupby(COL_CATEGORY)[COL_AMOUNT].agg(["sum", "count"]).reindex(order).dropna()
    g = g[g["count"] > 0]
    total = g["sum"].sum()
    g["비중(%)"] = (g["sum"] / total * 100).round(1) if total else 0.0
    g = g.rename(columns={"sum": "지출금액", "count": "건수"}).astype({"지출금액": "int64", "건수": "int64"})
    return g.sort_values("지출금액", ascending=False).reset_index()


# ── 엑셀 보고서 ──────────────────────────────────────────
SUMMARY_HEADER_COLOR = "#3d4451"  # 카테고리별 합계·전체내역 표의 머리글 색
EXCEL_TOP, EXCEL_LEFT = 2, 2      # 모든 시트는 1행·1열을 비우고 B2부터 시작
BLOCK_GAP = 2                     # 표와 표 사이에 비워 둘 행 수
AMOUNT_FORMAT = '#,##0"원"'


def _text_width(value) -> int:
    """엑셀 열 너비 계산용: 한글 등 전각 문자는 2칸으로 계산"""
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in str(value if value is not None else ""))


def _readable_font_color(hex_color: str) -> str:
    """머리글 배경색 위에서 잘 보이는 글자색(흰색/검정)"""
    def luminance(h: str) -> float:  # WCAG 상대 휘도
        ch = [int(h[i:i + 2], 16) / 255 for i in (0, 2, 4)]
        ch = [c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4 for c in ch]
        return 0.2126 * ch[0] + 0.7152 * ch[1] + 0.0722 * ch[2]

    bg = luminance(hex_color.lstrip("#"))
    white_contrast = 1.05 / (bg + 0.05)
    dark_contrast = (bg + 0.05) / (luminance("1A1A1A") + 0.05)
    return "FFFFFF" if white_contrast >= dark_contrast else "1A1A1A"


def _sheet_title(name: str, used: set[str]) -> str:
    """엑셀 시트 이름 규칙(31자, / \\ ? * [ ] : 금지)에 맞게 변환"""
    title = re.sub(r"[\\/?*\[\]:]", "·", name)[:31] or "시트"
    base, n = title, 2
    while title in used:
        title = f"{base[:28]}_{n}"
        n += 1
    used.add(title)
    return title


class _ExcelReport:
    """표(Excel Table) 단위로 시트를 채우고 열 너비를 맞춰 주는 도우미"""

    def __init__(self):
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.used_titles: set[str] = set()
        self.table_no = 0
        self.widths: dict[str, dict[int, int]] = {}

    def sheet(self, name: str, tab_color: Optional[str] = None):
        ws = self.wb.create_sheet(_sheet_title(name, self.used_titles))
        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = 2  # 들여쓰기용 빈 열
        if tab_color:
            ws.sheet_properties.tabColor = tab_color.lstrip("#").upper()
        self.widths[ws.title] = {}
        return ws

    def _track(self, ws, col: int, value) -> None:
        w = self.widths[ws.title]
        w[col] = max(w.get(col, 0), _text_width(value))

    def title(self, ws, row: int, text: str, color: str) -> None:
        cell = ws.cell(row=row, column=EXCEL_LEFT, value=text)
        cell.font = Font(bold=True, size=12)
        cell.border = Border(left=Side(style="thick", color=color.lstrip("#").upper()))

    def table(self, ws, top: int, headers: list[str], rows: list[list], color: str, formats: dict[str, str]) -> int:
        """top 행에 머리글이 오는 표를 그리고, 표의 마지막 행 번호를 반환한다."""
        fill = PatternFill("solid", fgColor=color.lstrip("#").upper())
        font = Font(bold=True, color=_readable_font_color(color))
        for j, h in enumerate(headers):
            cell = ws.cell(row=top, column=EXCEL_LEFT + j, value=h)
            cell.fill, cell.font = fill, font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            self._track(ws, j, h)
        for i, values in enumerate(rows, start=1):
            for j, (h, v) in enumerate(zip(headers, values)):
                cell = ws.cell(row=top + i, column=EXCEL_LEFT + j, value=v)
                if h in formats:
                    cell.number_format = formats[h]
                shown = f"{v:,}원" if isinstance(v, (int, float)) and h in formats else v
                self._track(ws, j, shown)

        self.table_no += 1
        last_row = top + max(len(rows), 1)
        ref = f"{get_column_letter(EXCEL_LEFT)}{top}:{get_column_letter(EXCEL_LEFT + len(headers) - 1)}{last_row}"
        table = Table(displayName=f"Table{self.table_no}", ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        ws.add_table(table)
        return last_row

    def total_row(self, ws, row: int, values: list, formats: list[Optional[str]]) -> None:
        line = Side(style="medium", color="3D4451")
        for j, (v, fmt) in enumerate(zip(values, formats)):
            cell = ws.cell(row=row, column=EXCEL_LEFT + j, value=v)
            cell.font = Font(bold=True)
            cell.border = Border(top=line, bottom=line)
            if fmt:
                cell.number_format = fmt

    def finish(self) -> bytes:
        for ws in self.wb.worksheets:
            for j, w in self.widths[ws.title].items():
                ws.column_dimensions[get_column_letter(EXCEL_LEFT + j)].width = min(max(w + 4, 8), 60)
        buf = io.BytesIO()
        self.wb.save(buf)
        return buf.getvalue()


def to_excel_bytes(detail: pd.DataFrame, summary: pd.DataFrame, colors: dict[str, str]) -> bytes:
    """
    분류 결과 보고서
      1) '요약' 시트: B2에 카테고리별 합계 표 → 행을 띄우고 지출이 큰 분류 순으로 분류별 내역 표
      2) 분류별 시트: 해당 분류의 내역 표 (번호 오름차순)
      3) '전체내역' 시트: 모든 내역 표
    모든 표는 엑셀 '표(Table)' 형식이며 머리글 행은 분류 색상으로 칠한다.
    """
    def rows_of(df: pd.DataFrame, with_category: bool) -> list[list]:
        df = df.sort_values(COL_NO)
        out = []
        for no, date, merchant, amount, cat in zip(df[COL_NO], df[COL_DATE], df[COL_MERCHANT], df[COL_AMOUNT], df[COL_CATEGORY]):
            row = [int(no), date.date() if pd.notna(date) else None, merchant, int(amount)]
            out.append(row + [cat] if with_category else row)
        return out

    detail_headers = [COL_NO, COL_DATE, COL_MERCHANT, COL_AMOUNT]
    detail_formats = {COL_DATE: "yyyy-mm-dd", COL_AMOUNT: AMOUNT_FORMAT}
    total_amount, total_count = int(summary["지출금액"].sum()), int(summary["건수"].sum())

    rpt = _ExcelReport()

    # 1) 요약 시트 ─ 카테고리별 합계
    ws = rpt.sheet("요약")
    summary_rows = [
        [r[COL_CATEGORY], int(r["지출금액"]), int(r["건수"]), float(r["비중(%)"]) / 100]
        for _, r in summary.iterrows()
    ]
    last = rpt.table(ws, EXCEL_TOP, [COL_CATEGORY, "지출금액", "건수", "비중"], summary_rows, SUMMARY_HEADER_COLOR,
                     {"지출금액": AMOUNT_FORMAT, "건수": '#,##0"건"', "비중": "0.0%"})
    rpt.total_row(ws, last + 1, ["합계", total_amount, total_count, 1.0 if total_count else 0.0],
                  [None, AMOUNT_FORMAT, '#,##0"건"', "0.0%"])
    # 분류명 칸 왼쪽에 분류 색상 띠를 둘러 차트 색과 연결
    for i, row in enumerate(summary_rows, start=1):
        ws.cell(row=EXCEL_TOP + i, column=EXCEL_LEFT).border = Border(
            left=Side(style="thick", color=colors.get(row[0], "#9a9993").lstrip("#").upper()))
    # '건수' 열은 아래 내역 표의 넓은 '가맹점' 열과 겹치므로 가운데 정렬
    for r in range(EXCEL_TOP + 1, last + 2):
        ws.cell(row=r, column=EXCEL_LEFT + 2).alignment = Alignment(horizontal="center")

    # 요약 시트 ─ 지출이 큰 분류 순으로 내역 표
    row = last + 1 + BLOCK_GAP + 1
    for cat, amount, count, _ in summary_rows:
        color = colors.get(cat, "#9a9993")
        rpt.title(ws, row, f"{cat}  ·  {count:,}건  ·  {amount:,}원", color)
        last = rpt.table(ws, row + 1, detail_headers, rows_of(detail[detail[COL_CATEGORY] == cat], False),
                         color, detail_formats)
        row = last + 1 + BLOCK_GAP

    # 2) 분류별 시트
    for cat, amount, count, _ in summary_rows:
        color = colors.get(cat, "#9a9993")
        ws = rpt.sheet(cat, tab_color=color)
        last = rpt.table(ws, EXCEL_TOP, detail_headers, rows_of(detail[detail[COL_CATEGORY] == cat], False),
                         color, detail_formats)
        rpt.total_row(ws, last + 1, ["합계", None, f"{count:,}건", amount], [None, None, None, AMOUNT_FORMAT])
        ws.freeze_panes = ws.cell(row=EXCEL_TOP + 1, column=1)

    # 3) 전체내역 시트
    ws = rpt.sheet("전체내역")
    last = rpt.table(ws, EXCEL_TOP, detail_headers + [COL_CATEGORY], rows_of(detail, True),
                     SUMMARY_HEADER_COLOR, detail_formats)
    rpt.total_row(ws, last + 1, ["합계", None, f"{total_count:,}건", total_amount, None],
                  [None, None, None, AMOUNT_FORMAT, None])
    ws.freeze_panes = ws.cell(row=EXCEL_TOP + 1, column=1)

    return rpt.finish()


def make_donut(summary: pd.DataFrame, colors: dict[str, str], order: list[str]):
    fig = px.pie(
        summary,
        names=COL_CATEGORY,
        values="지출금액",
        hole=0.6,
        color=COL_CATEGORY,
        color_discrete_map=colors,
        category_orders={COL_CATEGORY: order},
    )
    fig.update_traces(
        sort=False,
        direction="clockwise",
        textposition="outside",
        texttemplate="%{label}<br>%{percent:.1%}",
        hovertemplate="<b>%{label}</b><br>%{value:,.0f}원<br>%{percent:.1%}<extra></extra>",
        marker=dict(line=dict(color="rgba(255,255,255,1)", width=2)),
    )
    total = int(summary["지출금액"].sum())
    fig.update_layout(
        showlegend=True,
        legend=dict(orientation="h", yanchor="top", y=-0.08, xanchor="center", x=0.5),
        margin=dict(t=30, b=30, l=10, r=10),
        height=440,
        annotations=[dict(text=f"총 지출<br><b>{total:,}원</b>", x=0.5, y=0.5, showarrow=False, font=dict(size=16))],
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
    )
    return fig


def render_keyword_manager(settings: dict) -> None:
    """사이드바: 기본 키워드 보기 + 내 키워드 추가/삭제 + 내 분류(색상) 관리"""
    st.header("🏷️ 분류 키워드")
    st.caption("내가 추가한 키워드는 기본 키워드보다 먼저 적용되고, 대시보드에 바로 반영돼요. "
               "가맹점명에 키워드가 포함되면 해당 분류로 들어가요.")

    flash = st.session_state.pop("flash", None)
    if flash:
        getattr(st, flash[0])(flash[1])

    user_cats = set(user_category_names(settings))
    colors = category_colors(settings)
    for cat in category_order(settings)[:-1]:  # '기타' 제외
        my_kws = settings["keywords"].get(cat, [])
        badge = "🆕 " if cat in user_cats else ""
        # key 고정: 키워드 수(라벨)가 바뀌어도 펼친 상태 유지
        with st.expander(f"{badge}{cat} · 내 키워드 {len(my_kws)}개", key=f"exp_{cat}"):
            if cat in user_cats:
                c1, c2 = st.columns([1, 1], vertical_alignment="bottom")
                c1.color_picker("색상", value=colors[cat], key=f"color_{cat}",
                                on_change=on_color_change, args=(cat,))
                c2.button("🗑️ 분류 삭제", key=f"del_{cat}", on_click=on_delete_category, args=(cat,),
                          width="stretch")
            st.multiselect(
                "내 키워드",
                options=my_kws,
                default=my_kws,
                key=f"kw_{cat}",
                accept_new_options=True,
                placeholder="키워드 입력 후 Enter",
                on_change=on_keywords_change,
                args=(cat,),
                help="입력하고 Enter를 누르면 추가, ✕를 누르면 삭제돼요.",
            )
            if cat in CATEGORY_KEYWORDS:
                st.caption("기본 키워드: " + ", ".join(CATEGORY_KEYWORDS[cat]))

    with st.form("new_category", clear_on_submit=True, border=True):
        st.markdown("**➕ 새 분류 만들기**")
        st.text_input("분류 이름", key="new_cat_name", placeholder="예: 구독, 경조사, 교육")
        st.color_picker("색상", value=next_suggested_color(settings), key="new_cat_color")
        st.text_input("키워드 (쉼표로 구분)", key="new_cat_keywords", placeholder="예: 넷플릭스, 유튜브프리미엄")
        st.form_submit_button("분류 추가", on_click=on_add_category, width="stretch")


# ─────────────────────────────────────────────────────────────
# 4. UI
# ─────────────────────────────────────────────────────────────

st.set_page_config(page_title="Credit Card Management", page_icon="🧾", layout="wide")

st.markdown(
    """
    <style>
      .block-container { padding-top: 2rem; max-width: 1200px; }
      div[data-testid="stMetric"] {
          border: 1px solid rgba(128,128,128,0.25);
          border-radius: 12px; padding: 14px 18px;
      }
      .step { font-size: 0.85rem; font-weight: 600; opacity: 0.6; letter-spacing: 0.04em; }
    </style>
    """,
    unsafe_allow_html=True,
)

settings = get_settings()
ORDER = category_order(settings)
COLORS = category_colors(settings)

st.title("🧾Credit Card Management")
st.caption("카드사 엑셀 명세서를 올리면 사용처를 자동으로 분류하고 한눈에 보여드려요. "
           "업로드한 파일은 서버에 저장되지 않으며 브라우저를 닫으면 사라집니다.")

with st.sidebar:
    st.header("사용 방법")
    st.markdown(
        "1. 카드사 홈페이지/앱에서 **이용내역을 엑셀로 다운로드**\n"
        "2. **파일 업로드**에 그대로 올리기\n"
        "3. 표에서 **분류**를 필요하면 수정\n"
        "4. **대시보드**에서 분류를 클릭해 상세 내역 확인\n"
        "5. 결과 엑셀 다운로드"
    )
    st.divider()
    render_keyword_manager(settings)

# ── STEP 1 업로드 ─────────────────────────────────────────
st.markdown('<div class="step">STEP 1</div>', unsafe_allow_html=True)
st.subheader("파일 업로드")
uploaded = st.file_uploader("카드 명세서 엑셀 파일 (.xlsx, .xls)", type=["xlsx", "xls"])

if uploaded is None:
    st.info("👆 카드사에서 받은 엑셀 파일을 그대로 업로드해 주세요.")
    st.stop()

file_bytes = uploaded.getvalue()
file_id = hashlib.md5(file_bytes).hexdigest()[:12]
try:
    sheets = read_raw_excel(file_bytes, uploaded.name)
except ImportError as e:
    st.error(f"필요한 패키지가 없어요: `{e.name}`. `pip install -r requirements.txt` 후 다시 실행해 주세요.")
    st.stop()
except Exception as e:  # 손상된 파일, 암호 걸린 파일 등
    st.error(f"엑셀 파일을 읽지 못했어요. 암호가 걸려 있거나 손상된 파일인지 확인해 주세요.\n\n오류: {e}")
    st.stop()

sheets = {name: df for name, df in sheets.items() if not df.dropna(how="all").empty}
located = locate_statement(sheets) if sheets else None

with st.expander("원본 데이터 미리보기", expanded=False):
    if located:
        st.dataframe(located[1].head(50), width="stretch", hide_index=True)
    elif sheets:
        st.dataframe(next(iter(sheets.values())).head(50), width="stretch", hide_index=True)

if located is None:
    st.error("명세서에서 **가맹점** 열과 **결제금액** 열을 찾지 못했어요. "
             "'원본 데이터 미리보기'에서 열 이름을 확인하고, 열 이름이 다르면 알려주세요.")
    st.stop()

_, df_src, found = located
_fmt = detect_file_format(file_bytes)
notes = [f"인식된 열 — 이용일자: `{found['date'] or '없음'}` · 가맹점: `{found['merchant']}` · 결제금액: `{found['amount']}`"]
if _fmt in ("html", "xml"):
    notes.append(f"확장자는 엑셀이지만 실제로는 {_fmt.upper()} 형식이라 자동 변환했어요.")
st.caption("ℹ️ " + "  \n".join(notes))

all_tx = build_transactions(df_src, found["merchant"], found["amount"], found["date"])
if all_tx.empty:
    st.error("결제금액을 숫자로 읽을 수 있는 행이 없어요.")
    st.stop()

# 자동 분류 (사용자 키워드 → 기본 키워드 순). 키워드가 바뀌면 매 실행마다 즉시 다시 분류된다.
user_rules = build_user_rules(settings)
all_tx[COL_CATEGORY] = all_tx[COL_MERCHANT].map(lambda n: classify_merchant(n, user_rules))

opt1, opt2 = st.columns(2)
with opt1:
    exclude_negative = st.toggle("취소·환불(마이너스 금액) 제외", value=False,
                                 help="켜면 마이너스 금액과 결제금액 0원(결제 전 취소 등) 내역을 분석에서 뺍니다. "
                                      "끄면 모든 내역이 포함되고 취소 금액은 지출에서 차감됩니다.")
# 기본은 전체 포함. 토글을 켰을 때만 0원 이하(취소·환불) 내역을 제외한다.
cancel_mask = all_tx[COL_AMOUNT] <= 0
transactions = all_tx[~cancel_mask] if exclude_negative else all_tx

picked_months: list[str] = []
if transactions[COL_DATE].notna().any():
    months = sorted(transactions[COL_DATE].dropna().dt.to_period("M").astype(str).unique())
    if len(months) > 1:
        with opt2:
            picked_months = st.multiselect("기간(월) 필터", months, default=months)
        transactions = transactions[transactions[COL_DATE].dt.to_period("M").astype(str).isin(picked_months)]
transactions = transactions.reset_index(drop=True)

if exclude_negative and cancel_mask.any():
    st.caption(f"취소·환불(마이너스 금액, 결제금액 0원) {int(cancel_mask.sum())}건을 분석에서 제외했어요.")

# 수동 분류는 행 고유번호(_rid) 기준으로 저장 → 필터/키워드가 바뀌어도 유지된다
overrides: dict[int, str] = st.session_state.setdefault(f"overrides_{file_id}", {})
overrides = {rid: cat for rid, cat in overrides.items() if cat in ORDER}  # 삭제된 분류는 무시
st.session_state[f"overrides_{file_id}"] = overrides


def apply_overrides(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out[COL_CATEGORY] = [overrides.get(rid, cat) for rid, cat in zip(out[COL_RID], out[COL_CATEGORY])]
    out.insert(0, COL_NO, range(1, len(out) + 1))
    return out


# ── STEP 2 분류 편집 ───────────────────────────────────────
st.markdown('<div class="step">STEP 2</div>', unsafe_allow_html=True)
st.subheader("분류 확인 및 수정")

# 방금 표에서 바꾼 내용을 먼저 반영해 둔다 → 아래의 건수/목록이 한 박자 늦지 않게
_prev = st.session_state.get("_last_editor_view")  # (편집기 key, 그때 보이던 행의 _rid 목록)
if _prev and isinstance(st.session_state.get(_prev[0]), dict):
    for pos, changes in st.session_state[_prev[0]].get("edited_rows", {}).items():
        pos = int(pos)
        if COL_CATEGORY in changes and pos < len(_prev[1]) and changes[COL_CATEGORY] in ORDER:
            overrides[_prev[1][pos]] = changes[COL_CATEGORY]

current = apply_overrides(transactions)
other_count =int((current[COL_CATEGORY] == DEFAULT_CATEGORY).sum())
st.caption(f"총 {len(current):,}건 중 {len(current) - other_count:,}건이 분류되었어요. "
           "'분류' 칸을 클릭하면 드롭다운으로 바꿀 수 있어요. (번호·이용일자·가맹점·결제금액은 수정할 수 없어요)  \n"
           "같은 가맹점이 자주 나온다면 왼쪽 **🏷️ 분류 키워드**에 등록해 두면 다음부터 자동으로 분류돼요.")

snapshot_key = f"other_snapshot_{file_id}"
f1, f2 = st.columns([3, 1], vertical_alignment="center")
with f1:
    # key 고정: 라벨의 건수가 바뀌어도 체크 상태가 풀리지 않게
    only_other = st.checkbox(f"'{DEFAULT_CATEGORY}'로 분류된 항목만 보기 ({other_count}건 남음)", value=False,
                             key=f"only_other_{file_id}")

if only_other:
    # 체크한 순간의 '기타' 목록을 고정해 둔다. 분류를 바꿔도 행이 사라지지 않아
    # 여러 건을 연달아 바꿀 수 있고, 바꾼 내용은 대시보드에 바로 반영된다.
    if snapshot_key not in st.session_state:
        st.session_state[snapshot_key] = list(current.loc[current[COL_CATEGORY] == DEFAULT_CATEGORY, COL_RID])
    view = current[current[COL_RID].isin(st.session_state[snapshot_key])]
    done_count = int((view[COL_CATEGORY] != DEFAULT_CATEGORY).sum())
    with f2:
        st.button(
            "🔄 분류한 항목 숨기기",
            width="stretch",
            disabled=done_count == 0,
            on_click=lambda: st.session_state.pop(snapshot_key, None),
            help="다른 분류로 바꾼 행을 목록에서 정리하고, 남은 '기타' 항목만 다시 보여줘요.",
        )
    if done_count:
        st.caption(f"✅ 이 목록에서 {done_count}건을 분류했어요. 계속 바꿔도 목록이 유지되고, "
                   "다 끝나면 **🔄 분류한 항목 숨기기**로 정리하세요.")
else:
    st.session_state.pop(snapshot_key, None)
    view = current

# 보이는 행 구성이 바뀌면 편집기를 새로 만든다 (이전 편집 내용은 overrides에 이미 저장됨)
editor_key = "editor_" + hashlib.md5(f"{file_id}|{list(view[COL_RID])}".encode()).hexdigest()[:12]
last_key = st.session_state.get("_last_editor_key")
if last_key and last_key != editor_key:
    st.session_state.pop(last_key, None)
st.session_state["_last_editor_key"] = editor_key
st.session_state["_last_editor_view"] = (editor_key, [int(r) for r in view[COL_RID]])

edited_view = st.data_editor(
    view,
    key=editor_key,
    width="stretch",
    hide_index=True,
    num_rows="fixed",
    height=min(420, 38 + 35 * max(len(view), 1)),
    column_order=[COL_NO, COL_DATE, COL_MERCHANT, COL_AMOUNT, COL_CATEGORY],
    column_config={
        COL_NO: st.column_config.NumberColumn(COL_NO, width="small"),
        COL_DATE: st.column_config.DateColumn(COL_DATE, format="YYYY-MM-DD"),
        COL_MERCHANT: st.column_config.TextColumn(COL_MERCHANT, width="large"),
        COL_AMOUNT: st.column_config.NumberColumn(COL_AMOUNT, format="%d원"),
        COL_CATEGORY: st.column_config.SelectboxColumn(
            COL_CATEGORY, options=ORDER, required=True, width="medium",
            help="클릭해서 분류를 변경하세요",
        ),
    },
    disabled=[COL_NO, COL_DATE, COL_MERCHANT, COL_AMOUNT],
)

# 편집 결과 반영: 자동 분류와 다르거나 이미 수동 지정한 행은 overrides에 기록
auto_by_rid = dict(zip(transactions[COL_RID], transactions[COL_CATEGORY]))
for rid, cat in zip(edited_view[COL_RID], edited_view[COL_CATEGORY]):
    rid = int(rid)
    if cat != auto_by_rid.get(rid) or rid in overrides:
        overrides[rid] = cat
final = apply_overrides(transactions)

changed = sum(1 for rid, c in overrides.items() if rid in auto_by_rid and c != auto_by_rid[rid])
if changed:
    c1, c2 = st.columns([3, 1], vertical_alignment="center")
    c1.caption(f"✏️ 직접 수정한 항목: {changed}건 (키워드보다 우선 적용돼요)")
    if c2.button("수정 내용 초기화", width="stretch"):
        st.session_state[f"overrides_{file_id}"] = {}
        st.session_state.pop(editor_key, None)
        st.rerun()

# ── STEP 3 대시보드 ───────────────────────────────────────
st.markdown('<div class="step">STEP 3</div>', unsafe_allow_html=True)
st.subheader("지출 대시보드")

summary = summarize(final, ORDER)
total = int(final[COL_AMOUNT].sum())
top = summary.iloc[0] if not summary.empty else None

k1, k2, k3 = st.columns(3)
k1.metric("총 지출 금액", f"{total:,}원")
k2.metric("이용 건수", f"{len(final):,}건")
k3.metric("최대 지출 분류", f"{top[COL_CATEGORY]}" if top is not None else "-",
          f"{int(top['지출금액']):,}원 · {top['비중(%)']}%" if top is not None else None, delta_color="off")

if final[COL_DATE].notna().any():
    d = final[COL_DATE].dropna()
    st.caption(f"기간: {d.min():%Y-%m-%d} ~ {d.max():%Y-%m-%d}")

left, right = st.columns([1.15, 1])
with left:
    st.markdown("**카테고리별 소비 비중**")
    chart_data = summary[summary["지출금액"] > 0]  # 0원뿐인 분류는 차트 조각이 없으므로 제외
    if summary["지출금액"].lt(0).any() or chart_data.empty:
        st.info("지출 합계가 마이너스인 분류가 있어 차트를 그릴 수 없어요. '취소·환불 제외'를 켜 보세요.")
    else:
        st.plotly_chart(make_donut(chart_data, COLORS, ORDER), width="stretch", config={"displayModeBar": False})
with right:
    st.markdown("**카테고리별 지출 요약**")
    st.caption("행을 클릭하면 아래에 해당 분류의 상세 내역이 나와요.")
    summary_view = summary.copy()
    summary_view.insert(0, "색", "●")
    selection = st.dataframe(
        summary_view.style.apply(
            lambda col: [f"color: {COLORS.get(c, '#9a9993')}" for c in summary_view[COL_CATEGORY]],
            subset=["색"],
        ),
        hide_index=True,
        width="stretch",
        on_select="rerun",
        selection_mode="single-row",
        key=f"summary_{file_id}",
        column_config={
            "색": st.column_config.TextColumn("", width=30),
            COL_CATEGORY: st.column_config.TextColumn("분류"),
            "지출금액": st.column_config.NumberColumn("지출금액", format="%d원"),
            "건수": st.column_config.NumberColumn("건수", format="%d건"),
            "비중(%)": st.column_config.ProgressColumn("비중", format="%.1f%%", min_value=0, max_value=100),
        },
    )

selected_rows = selection.selection.rows if selection else []
if selected_rows and selected_rows[0] < len(summary):
    sel_cat = summary.iloc[selected_rows[0]][COL_CATEGORY]
    detail = final[final[COL_CATEGORY] == sel_cat].sort_values(COL_DATE, na_position="last")
    st.markdown(f"#### 📂 {sel_cat} 상세 내역 · {len(detail):,}건 · {int(detail[COL_AMOUNT].sum()):,}원")
    st.dataframe(
        detail,
        hide_index=True,
        width="stretch",
        column_order=[COL_NO, COL_DATE, COL_MERCHANT, COL_AMOUNT],
        column_config={
            COL_NO: st.column_config.NumberColumn(COL_NO, width="small"),
            COL_DATE: st.column_config.DateColumn(COL_DATE, format="YYYY-MM-DD"),
            COL_MERCHANT: st.column_config.TextColumn(COL_MERCHANT, width="large"),
            COL_AMOUNT: st.column_config.NumberColumn(COL_AMOUNT, format="%d원"),
        },
    )

# ── 다운로드 ──────────────────────────────────────────────
st.divider()
base_name = re.sub(r"\.xlsx?$", "", uploaded.name, flags=re.IGNORECASE)
st.download_button(
    "📥 분류 결과 엑셀 다운로드",
    data=to_excel_bytes(final, summary, COLORS),
    file_name=f"{base_name}_분류결과.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    type="primary",
    width="stretch",
)
st.caption("다운로드 파일 구성: '요약' 시트(카테고리별 합계 + 지출이 큰 분류 순 내역) · 분류별 시트 · '전체내역' 시트")
